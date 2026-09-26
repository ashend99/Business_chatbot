"""Chunking and file/URL text extraction for the documents pipeline."""

from io import BytesIO

import httpx
import pytest

from app.services import chunking, file_parsing
from app.services.chunking import approx_token_count, chunk_text

# ---- chunking ---------------------------------------------------------------


def test_chunk_empty_or_whitespace_text() -> None:
    assert chunk_text("") == []
    assert chunk_text("   \n\n  ") == []


def test_short_text_is_one_chunk() -> None:
    assert chunk_text("  Opening hours are 9 to 5.  ") == ["Opening hours are 9 to 5."]


def test_long_text_is_split_within_target_size_with_overlap() -> None:
    sentences = [f"Sentence number {i} talks about topic {i % 7}." for i in range(400)]
    chunks = chunk_text(" ".join(sentences), target_tokens=100, overlap_tokens=20)
    max_chars = 100 * chunking.CHARS_PER_TOKEN
    assert len(chunks) > 1
    assert all(len(c) <= max_chars for c in chunks)
    # overlap: consecutive chunks share some text across the boundary
    assert any(a[-40:].split()[-1] in b for a, b in zip(chunks, chunks[1:], strict=False))
    # nothing lost
    assert "Sentence number 399" in chunks[-1]


def test_paragraph_boundaries_are_preferred() -> None:
    paragraph = "word " * 300
    chunks = chunk_text(f"{paragraph}\n\n{paragraph}", target_tokens=100, overlap_tokens=0)
    assert all(not c.startswith(" ") for c in chunks)


def test_approx_token_count() -> None:
    assert approx_token_count("") == 1  # never zero
    assert approx_token_count("a" * 400) == 100


# ---- file parsing -----------------------------------------------------------


def test_parse_txt_handles_bad_utf8() -> None:
    assert file_parsing.parse_txt(b"  hello \xff world  ") == "hello � world"


def test_parse_docx() -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Refund policy")
    doc.add_paragraph("   ")
    doc.add_paragraph("Returns accepted within 7 days.")
    buffer = BytesIO()
    doc.save(buffer)
    assert (
        file_parsing.parse_docx(buffer.getvalue())
        == "Refund policy\n\nReturns accepted within 7 days."
    )


def test_parse_pdf() -> None:
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, "Delivery takes two days.")
    pdf.showPage()
    pdf.drawString(72, 720, "Pickup is free.")
    pdf.save()
    text = file_parsing.parse_pdf(buffer.getvalue())
    assert "Delivery takes two days." in text
    assert "Pickup is free." in text


async def test_fetch_url_text_strips_page_chrome(monkeypatch) -> None:
    html = """
    <html><head><style>body{}</style><script>track()</script></head>
    <body><nav>Home | About</nav><header>Logo</header>
    <main><h1>Our Story</h1><p>Family run   since 1990.</p></main>
    <footer>(c) 2026</footer></body></html>
    """
    seen = {}

    def handler(request: httpx.Request) -> httpx.Response:
        seen["url"] = str(request.url)
        return httpx.Response(200, text=html)

    real_client = httpx.AsyncClient
    monkeypatch.setattr(
        file_parsing.httpx,
        "AsyncClient",
        lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw),
    )
    text = await file_parsing.fetch_url_text("https://example.com/about")
    assert seen["url"] == "https://example.com/about"
    assert text == "Our Story\nFamily run   since 1990."


async def test_fetch_url_text_raises_on_http_error(monkeypatch) -> None:
    real_client = httpx.AsyncClient
    monkeypatch.setattr(
        file_parsing.httpx,
        "AsyncClient",
        lambda **kw: real_client(
            transport=httpx.MockTransport(lambda r: httpx.Response(404)), **kw
        ),
    )
    with pytest.raises(httpx.HTTPStatusError):
        await file_parsing.fetch_url_text("https://example.com/missing")
